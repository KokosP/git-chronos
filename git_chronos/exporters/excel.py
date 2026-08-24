# Copyright (c) sbgid.com
"""Corporate-style Excel invoice generation via openpyxl."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..analyzer import AnalysisResult
from ..config import Config

HEADER_COLOR = "34495E"
ALT_ROW_COLOR = "ECF0F1"
DATE_FMT = "YYYY-MM-DD"
HOURS_FMT = "0.0"

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(size=16, bold=True, color=HEADER_COLOR)
_LABEL_FONT = Font(bold=True)
_TOTAL_FONT = Font(bold=True)
_THIN_BOTTOM = Border(bottom=Side(style="thin", color="D5DBDB"))


def _money_fmt(symbol: str) -> str:
    return f'"{symbol}"#,##0.00'


def export_excel(
    result: AnalysisResult,
    config: Config,
    out_dir: Path,
    period_label: str,
) -> Path:
    """Write the report to an .xlsx file and return its path."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor=HEADER_COLOR)
    alt_fill = PatternFill("solid", fgColor=ALT_ROW_COLOR)
    money_fmt = _money_fmt(config.currency_symbol)

    ws["A1"] = config.company_name
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = "TIME REPORT / INVOICE"
    ws["A2"].font = Font(color=HEADER_COLOR, bold=True)

    meta = [
        ("Developer:", f"{config.developer_name} <{config.developer_email}>"),
        ("Client:", config.client_name),
        ("Period:", period_label),
        ("Generated:", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Rate:", f"{config.currency_symbol}{config.hourly_rate:.2f}/h"),
    ]
    for i, (label, value) in enumerate(meta, start=4):
        ws.cell(row=i, column=1, value=label).font = _LABEL_FONT
        ws.cell(row=i, column=2, value=value)

    header_row = 10
    headers = ["Date", "Sessions", "Hours", "Rate", "Amount", "Tasks"]
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.font = _HEADER_FONT
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    first_data_row = header_row + 1
    for i, day in enumerate(result.days):
        row = first_data_row + i
        ws.cell(row=row, column=1, value=day.day).number_format = DATE_FMT
        ws.cell(row=row, column=2, value=len(day.sessions))
        ws.cell(row=row, column=3, value=round(day.hours, 2)).number_format = HOURS_FMT
        ws.cell(row=row, column=4, value=result.hourly_rate).number_format = money_fmt
        ws.cell(row=row, column=5, value=f"=C{row}*D{row}").number_format = money_fmt
        ws.cell(row=row, column=6, value=", ".join(day.tasks))
        if i % 2 == 1:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row, column=col).fill = alt_fill
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).border = _THIN_BOTTOM

    last_data_row = first_data_row + len(result.days) - 1
    total_row = last_data_row + 1
    ws.cell(row=total_row, column=1, value="TOTAL").font = _TOTAL_FONT
    ws.cell(row=total_row, column=3,
            value=f"=SUM(C{first_data_row}:C{last_data_row})").number_format = HOURS_FMT
    ws.cell(row=total_row, column=5,
            value=f"=SUM(E{first_data_row}:E{last_data_row})").number_format = money_fmt
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).font = _TOTAL_FONT
        ws.cell(row=total_row, column=col).border = Border(top=Side(style="double"))

    for col, width in zip(range(1, len(headers) + 1), (14, 10, 10, 12, 14, 32)):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f"A{first_data_row}"

    filename = f"git-chronos-invoice-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    out_path = Path(out_dir) / filename
    wb.save(out_path)
    return out_path
